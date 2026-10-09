from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "documents" / "Monografia_SmartFix_Atualizada_2026.docx"
ASSETS = ROOT / "tmp" / "monografia_assets"
OUT.parent.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

AUTHORS = [
    "ARTHUR FORTUNATO DA SILVA",
    "GABRIEL SILVA FOGAÇA",
    "GEAN DA SILVA ABREU",
    "FLÁVIO HENRIQUE PRADO",
]

TITLE = "SMARTFIX"
SUBTITLE = "Engenharia de software aplicada ao desenvolvimento de uma plataforma web para gerenciamento e acompanhamento de reparos de dispositivos eletrônicos"


def font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_table_borders(table, color="B7B7B7", size="4"):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(3)
section.left_margin = Cm(3)
section.right_margin = Cm(2)
section.bottom_margin = Cm(2)
section.page_width = Cm(21)
section.page_height = Cm(29.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.space_after = Pt(0)

for name, size in (("Title", 14), ("Heading 1", 12), ("Heading 2", 12), ("Heading 3", 12), ("Heading 4", 12)):
    st = styles[name]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
    if name == "Heading 1":
        st.paragraph_format.page_break_before = True

if "Caption" not in styles:
    styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
cap = styles["Caption"]
cap.font.name = "Times New Roman"
cap.font.size = Pt(10)
cap.font.color.rgb = RGBColor(0, 0, 0)
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.line_spacing = 1
cap.paragraph_format.space_after = Pt(4)


def add_plain(text="", align=None, indent=True, bold=False, italic=False, size=12, spacing=1.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    r = p.add_run(text)
    font(r, size=size, bold=bold, italic=italic)
    return p


def add_heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text)
    font(r, 12, bold=True)
    return p


def add_table(headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1
        r = p.add_run(str(h))
        font(r, font_size, bold=True)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx % 2:
                set_cell_shading(cell, "EAF2F8")
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(str(value)) > 18 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            font(r, font_size)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_caption(title, source="Elaborado pelos autores (2026)."):
    p = doc.add_paragraph(style="Caption")
    r = p.add_run(title)
    font(r, 10, bold=True)
    p2 = doc.add_paragraph(style="Caption")
    r2 = p2.add_run(f"Fonte: {source}")
    font(r2, 10)


def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(item)
        font(r)


def page_break():
    doc.add_page_break()


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def draw_diagram(path, title, columns):
    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    try:
        f_title = ImageFont.truetype("arial.ttf", 38)
        f_box = ImageFont.truetype("arial.ttf", 25)
    except Exception:
        f_title = ImageFont.load_default()
        f_box = ImageFont.load_default()
    d.text((800, 45), title, fill="black", font=f_title, anchor="mm")
    colw = 1450 // len(columns)
    centers = []
    for ci, (label, boxes) in enumerate(columns):
        x0 = 60 + ci * colw
        x1 = x0 + colw - 25
        d.rounded_rectangle((x0, 105, x1, 820), radius=18, outline="#4F81BD", width=4, fill="#F3F7FB")
        d.text(((x0+x1)//2, 135), label, fill="#1F4E78", font=f_box, anchor="mm")
        ys=[]
        step = 590 // max(len(boxes), 1)
        for bi, box in enumerate(boxes):
            y = 205 + bi * step
            d.rounded_rectangle((x0+30, y, x1-30, y+85), radius=12, outline="#6B6B6B", width=3, fill="white")
            d.multiline_text(((x0+x1)//2, y+42), box, fill="black", font=f_box, anchor="mm", align="center", spacing=4)
            ys.append((x0+30, x1-30, y+42))
        centers.append(ys)
    for ci in range(len(centers)-1):
        if not centers[ci] or not centers[ci+1]:
            continue
        left = centers[ci][min(len(centers[ci])-1, len(centers[ci])//2)]
        right = centers[ci+1][min(len(centers[ci+1])-1, len(centers[ci+1])//2)]
        x0, y0 = left[1], left[2]
        x1, y1 = right[0], right[2]
        d.line((x0, y0, x1, y1), fill="#1F4E78", width=5)
        d.polygon([(x1, y1), (x1-18, y1-10), (x1-18, y1+10)], fill="#1F4E78")
    img.save(path)


def draw_states(path):
    img = Image.new("RGB", (1800, 700), "white")
    d = ImageDraw.Draw(img)
    try:
        ft = ImageFont.truetype("arial.ttf", 34)
        fb = ImageFont.truetype("arial.ttf", 23)
    except Exception:
        ft = fb = ImageFont.load_default()
    d.text((900, 45), "Ciclo da ordem de serviço", fill="black", font=ft, anchor="mm")
    labels = ["Aguardando\ndiagnóstico", "Aguardando\naprovação", "Orçamento\naprovado", "Em reparo", "Aguardando\npeças", "Pronto para\nretirada", "Concluído"]
    pos = [(90,170),(330,170),(570,170),(810,170),(1050,330),(1290,170),(1530,170)]
    for (x,y),label in zip(pos,labels):
        d.rounded_rectangle((x,y,x+190,y+100),radius=16,outline="#1F4E78",width=4,fill="#EAF2F8")
        d.multiline_text((x+95,y+50),label,font=fb,fill="black",anchor="mm",align="center")
    seq=[(0,1),(1,2),(2,3),(3,5),(5,6),(3,4),(4,3)]
    for a,b in seq:
        xa,ya=pos[a]; xb,yb=pos[b]
        x0=xa+190 if xb>xa else xa
        x1=xb if xb>xa else xb+190
        y0=ya+50; y1=yb+50
        d.line((x0,y0,x1,y1),fill="#1F4E78",width=4)
        if xb>xa: d.polygon([(x1,y1),(x1-15,y1-8),(x1-15,y1+8)],fill="#1F4E78")
        else: d.polygon([(x1,y1),(x1+15,y1-8),(x1+15,y1+8)],fill="#1F4E78")
    d.rounded_rectangle((330,480,520,580),radius=16,outline="#8B0000",width=4,fill="#FDEDEC")
    d.text((425,530),"Cancelado",font=fb,fill="black",anchor="mm")
    d.line((425,270,425,480),fill="#8B0000",width=4)
    d.polygon([(425,480),(417,465),(433,465)],fill="#8B0000")
    img.save(path)


arch_img = ASSETS / "arquitetura.png"
use_img = ASSETS / "casos_uso.png"
der_img = ASSETS / "der.png"
state_img = ASSETS / "estados.png"
draw_diagram(arch_img, "Arquitetura lógica", [
    ("Usuários", ["Navegador desktop", "Navegador móvel"]),
    ("Aplicação Next.js", ["Páginas React", "Route Handlers", "Controllers", "Serviços e políticas"]),
    ("Persistência", ["Sequelize", "PostgreSQL"]),
    ("Integrações", ["Google OAuth", "Photon / OSM", "Resend"]),
])
draw_diagram(use_img, "Casos de uso por perfil", [
    ("Cliente", ["Endereços e dispositivos", "Buscar assistência", "Solicitar e acompanhar", "Avaliar e conversar"]),
    ("Parceiro", ["Catálogo de serviços", "Diagnosticar e orçar", "Atualizar reparo", "Notificações"]),
    ("Administrador", ["Aprovar parceiros", "Atendimento SmartFix"]),
])
draw_diagram(der_img, "Entidades principais", [
    ("Contas", ["Clientes", "Parceiros", "Endereços"]),
    ("Operação", ["Dispositivos", "Ordens de reparo", "Avaliações"]),
    ("Relacionamento", ["Serviços do parceiro", "Mensagens", "Workflow"]),
])
draw_states(state_img)


# Capa
for _ in range(2): add_plain("", indent=False)
add_plain("FACULDADE DE TECNOLOGIA PADRE DANILO JOSÉ DE OLIVEIRA OHL", WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12)
for a in AUTHORS:
    add_plain(a, WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12, spacing=1)
for _ in range(5): add_plain("", indent=False)
add_plain(TITLE, WD_ALIGN_PARAGRAPH.CENTER, False, True, size=14)
add_plain(SUBTITLE.upper(), WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12)
for _ in range(8): add_plain("", indent=False)
add_plain("BARUERI", WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12)
add_plain("2026", WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12)
page_break()

# Folha de rosto
for a in AUTHORS:
    add_plain(a, WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12, spacing=1)
for _ in range(4): add_plain("", indent=False)
add_plain(TITLE, WD_ALIGN_PARAGRAPH.CENTER, False, True, size=14)
add_plain(SUBTITLE, WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12)
for _ in range(3): add_plain("", indent=False)
p = add_plain("Monografia apresentada à Faculdade de Tecnologia Padre Danilo José de Oliveira Ohl como requisito para conclusão do curso de Análise e Desenvolvimento de Sistemas.", WD_ALIGN_PARAGRAPH.JUSTIFY, False, size=12)
p.paragraph_format.left_indent = Cm(8)
p.paragraph_format.line_spacing = 1
p2 = add_plain("Orientador: Prof. Vander Ribeiro Eime.", WD_ALIGN_PARAGRAPH.LEFT, False, size=12, spacing=1)
p2.paragraph_format.left_indent = Cm(8)
for _ in range(5): add_plain("", indent=False)
add_plain("BARUERI", WD_ALIGN_PARAGRAPH.CENTER, False, True)
add_plain("2026", WD_ALIGN_PARAGRAPH.CENTER, False, True)
page_break()

# Aprovação
add_plain("FOLHA DE APROVAÇÃO", WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12)
for a in AUTHORS:
    add_plain(a, WD_ALIGN_PARAGRAPH.CENTER, False, True, size=11, spacing=1)
add_plain(TITLE, WD_ALIGN_PARAGRAPH.CENTER, False, True, size=12)
add_plain(SUBTITLE, WD_ALIGN_PARAGRAPH.CENTER, False, False, size=12)
add_plain("Monografia aprovada como requisito para conclusão do curso de Análise e Desenvolvimento de Sistemas.", WD_ALIGN_PARAGRAPH.JUSTIFY, False)
add_plain("Data de aprovação: ____ de __________________ de 2026.", WD_ALIGN_PARAGRAPH.LEFT, False)
for label in ["Prof. Vander Ribeiro Eime - Orientador", "Professor(a) examinador(a)", "Professor(a) examinador(a)"]:
    add_plain("\n____________________________________________", WD_ALIGN_PARAGRAPH.CENTER, False, spacing=1)
    add_plain(label, WD_ALIGN_PARAGRAPH.CENTER, False, spacing=1)
page_break()

# Resumo
add_heading("RESUMO", 1)
resumo = (
    "O processo de manutenção de dispositivos eletrônicos envolve troca de informações entre clientes e assistências técnicas, incluindo identificação do equipamento, descrição do defeito, diagnóstico, orçamento e acompanhamento do serviço. Quando essas informações permanecem distribuídas em canais distintos, aumentam as dificuldades de organização e rastreabilidade. Este trabalho apresenta o desenvolvimento do SmartFix, uma plataforma web responsiva destinada a centralizar esse relacionamento. A pesquisa é aplicada, de caráter descritivo, e utiliza levantamento de requisitos, modelagem, desenvolvimento incremental e validação técnica. A solução foi implementada com Next.js, React, TypeScript, PostgreSQL, Sequelize e Zod. O sistema contempla cadastro e autenticação, endereços, dispositivos, descoberta geográfica de assistências aprovadas, solicitação e triagem, diagnóstico, orçamento, decisão do cliente, histórico da ordem, notificações, mensagens e avaliações. A validação executou 42 testes automatizados, todos aprovados, além de análise estática, build de produção e verificações transacionais no PostgreSQL. O banco confirmou Row Level Security nas nove tabelas avaliadas e bloqueio de acesso direto por papéis públicos. Os resultados demonstram a viabilidade da plataforma para organizar o ciclo de reparo dentro do escopo implementado. Pagamento integrado, logística e relatórios administrativos permanecem como evoluções futuras."
)
add_plain(resumo)
add_plain("Palavras-chave: Engenharia de software. Plataforma web. Assistência técnica. Ordem de serviço. Rastreabilidade.", indent=False, bold=True, spacing=1)
page_break()

add_heading("ABSTRACT", 1)
abstract = (
    "The repair of electronic devices requires information exchange between customers and technical service providers, including device identification, defect description, diagnosis, quotation, and service tracking. When this information remains distributed across different channels, organization and traceability become more difficult. This study presents the development of SmartFix, a responsive web platform designed to centralize this relationship. The work is an applied and descriptive study based on requirements analysis, system modeling, incremental development, and technical validation. The solution was implemented with Next.js, React, TypeScript, PostgreSQL, Sequelize, and Zod. It includes registration and authentication, addresses, devices, geographic discovery of approved service providers, repair requests, initial triage, diagnosis, quotations, customer decisions, order history, notifications, messaging, and reviews. Validation executed 42 automated tests, all successful, as well as static analysis, a production build, and transactional PostgreSQL checks. The database verification confirmed Row Level Security on the nine evaluated tables and blocked direct access by public roles. The results indicate that the platform can organize the repair lifecycle within the implemented scope. Integrated payments, logistics, and administrative reports remain future work."
)
add_plain(abstract)
add_plain("Keywords: Software engineering. Web platform. Technical assistance. Service order. Traceability.", indent=False, bold=True, spacing=1)
page_break()

# Sumário simples
add_heading("SUMÁRIO", 1)
toc_items = [
    "1 INTRODUÇÃO", "2 FUNDAMENTAÇÃO E CONTEXTO DO PROBLEMA", "3 METODOLOGIA",
    "4 ESPECIFICAÇÃO DO SMARTFIX", "5 MODELAGEM E ARQUITETURA",
    "6 DESENVOLVIMENTO E IMPLEMENTAÇÃO", "7 TESTES RESULTADOS E DISCUSSÃO",
    "8 CONSIDERAÇÕES FINAIS", "REFERÊNCIAS", "APÊNDICES"
]
for item in toc_items:
    add_plain(item, WD_ALIGN_PARAGRAPH.LEFT, False, size=12, spacing=1.5)

# Capítulo 1
add_heading("1 INTRODUÇÃO", 1)
for text in [
    "A presença de dispositivos eletrônicos nas atividades pessoais, acadêmicas e profissionais tornou a manutenção desses equipamentos um serviço recorrente. Apesar dessa demanda, a relação entre clientes e assistências técnicas ainda pode depender de contatos fragmentados, registros informais e pouca visibilidade sobre diagnóstico, orçamento e andamento do reparo. Nesse contexto, a informação deixa de acompanhar o serviço de maneira estruturada e pode gerar dúvidas, retrabalho e dificuldade de responsabilização.",
    "O SmartFix foi concebido como uma plataforma web capaz de organizar esse processo. A proposta não consiste em executar diretamente a manutenção, mas em oferecer uma infraestrutura digital para clientes e assistências parceiras registrarem e acompanharem as etapas relevantes. A solução reúne cadastro de equipamentos, descoberta de prestadores, abertura da solicitação, triagem, diagnóstico, orçamento, aprovação, evolução do serviço e avaliação final.",
    "A versão analisada representa uma mudança em relação ao projeto inicialmente documentado como aplicativo móvel. A implementação atual utiliza tecnologias web e pode ser acessada por navegadores em dispositivos móveis ou computadores. Essa mudança amplia o acesso sem exigir instalação por loja de aplicativos e aproxima a interface dos módulos de cliente, parceiro e administração.",
]: add_plain(text)
add_heading("1.1 Problema de pesquisa", 2)
add_plain("A pesquisa é orientada pela seguinte questão: como uma plataforma web pode centralizar e tornar mais transparente o relacionamento entre clientes e assistências técnicas durante o ciclo de reparo de dispositivos eletrônicos?")
add_heading("1.2 Justificativa", 2)
add_plain("A relevância do trabalho está na necessidade de preservar informações operacionais e tornar compreensíveis as responsabilidades de cada participante. Um histórico estruturado permite distinguir solicitação, diagnóstico, proposta comercial, decisão e execução. Para a assistência, a centralização reduz a dispersão de dados; para o cliente, oferece um ponto de consulta sobre o serviço. O projeto também constitui uma aplicação prática de requisitos, arquitetura web, modelagem relacional, segurança, testes e integração entre interface e servidor.")
add_heading("1.3 Objetivo geral", 2)
add_plain("Desenvolver e avaliar uma plataforma web responsiva para organizar a solicitação, o diagnóstico, o orçamento e o acompanhamento de reparos realizados por assistências técnicas parceiras.")
add_heading("1.4 Objetivos específicos", 2)
add_bullets([
    "levantar e formalizar requisitos e regras de negócio do processo de reparo;",
    "separar permissões de cliente, parceiro e administrador;",
    "modelar dispositivos, endereços, serviços, ordens, avaliações e mensagens;",
    "implementar busca geográfica de assistências aprovadas;",
    "controlar diagnóstico, orçamento e transições da ordem de serviço;",
    "preservar histórico e isolamento dos dados;",
    "validar a solução por testes, análise estática, build e verificações do banco."
])
add_heading("1.5 Delimitação", 2)
add_plain("O escopo entregue não inclui pagamento integrado, coleta, entrega, rastreamento logístico, moderação de avaliações ou relatórios administrativos completos. Google OAuth e recuperação por e-mail dependem de configuração externa. Essas condições são explicitadas para que recursos planejados não sejam confundidos com resultados alcançados.")
add_heading("1.6 Organização do trabalho", 2)
add_plain("Após esta introdução, o Capítulo 2 apresenta a fundamentação e o contexto. O Capítulo 3 descreve a metodologia. Os Capítulos 4 e 5 registram especificação, modelagem e arquitetura. O Capítulo 6 detalha a implementação e as interfaces. O Capítulo 7 apresenta a validação e discute os resultados. O Capítulo 8 reúne as considerações finais e os trabalhos futuros.")

# Capítulo 2
add_heading("2 FUNDAMENTAÇÃO E CONTEXTO DO PROBLEMA", 1)
add_heading("2.1 Engenharia de software e requisitos", 2)
add_plain("A engenharia de software reúne métodos e práticas para especificar, desenvolver, validar e evoluir sistemas. Em projetos com diferentes perfis de usuário, os requisitos precisam descrever não apenas funcionalidades, mas também propriedade dos dados, permissões e condições de transição. A rastreabilidade relaciona cada requisito à sua implementação e verificação, reduzindo divergências entre documentação e produto (SOMMERVILLE, 2019; PRESSMAN; MAXIM, 2021).")
add_heading("2.2 Plataformas web e separação de responsabilidades", 2)
add_plain("A aplicação utiliza o App Router do Next.js, cuja organização baseada em arquivos permite definir páginas, layouts e Route Handlers. Os Route Handlers processam métodos HTTP no diretório da aplicação e podem funcionar como camada de backend para a interface (NEXT.JS, 2026a; NEXT.JS, 2026b). No SmartFix, essa capacidade é combinada com controllers, serviços, validações e models para separar apresentação, regra de negócio e persistência.")
add_heading("2.3 Interfaces responsivas", 2)
add_plain("Uma interface responsiva adapta a organização visual ao espaço disponível sem alterar a finalidade do fluxo. Essa propriedade é relevante porque clientes podem iniciar e acompanhar solicitações pelo celular, enquanto assistências podem preferir telas maiores para consultar ordens e preencher orçamentos. A utilização de componentes React permite organizar a interface em unidades reutilizáveis e manter estado e interação de formulários (REACT, 2024).")
add_heading("2.4 Bancos relacionais e integridade", 2)
add_plain("O banco relacional é adequado quando o domínio exige vínculos verificáveis entre contas, dispositivos, ordens e avaliações. Chaves estrangeiras, unicidade e restrições permitem impedir combinações inválidas mesmo que uma requisição inadequada alcance a persistência. O PostgreSQL também oferece Row Level Security, mecanismo que restringe operações sobre linhas conforme políticas e assume negação por padrão quando a proteção está habilitada sem política aplicável (POSTGRESQL, 2026).")
add_heading("2.5 Proteção de dados", 2)
add_plain("A plataforma trata informações de pessoas identificadas, como nome, contato, documento e endereço. A Lei nº 13.709/2018 dispõe sobre o tratamento de dados pessoais, inclusive em meios digitais, com o objetivo de proteger liberdade, privacidade e desenvolvimento da personalidade (BRASIL, 2018). A implementação técnica não equivale, por si só, a uma certificação de conformidade, mas deve apoiar princípios como finalidade, necessidade, segurança e controle de acesso.")
add_heading("2.6 Localização e descoberta", 2)
add_plain("A descoberta de prestadores exige distinguir localização aproximada de rota. O SmartFix geocodifica uma origem e os endereços de parceiros aprovados, calcula distância geográfica e aplica um raio. O resultado ajuda na ordenação das alternativas, mas não representa tempo de viagem nem serviço logístico.")
add_heading("2.7 Síntese", 2)
add_plain("A fundamentação orienta quatro decisões do projeto: especificar regras verificáveis; separar interface, servidor e banco; reforçar a integridade em mais de uma camada; e apresentar ao usuário somente informações compatíveis com seu papel e propriedade.")

# Capítulo 3
add_heading("3 METODOLOGIA", 1)
add_heading("3.1 Natureza do trabalho", 2)
add_plain("O trabalho caracteriza-se como pesquisa aplicada, pois utiliza conhecimentos de engenharia de software para produzir uma solução destinada a um problema operacional. A abordagem é predominantemente qualitativa e descritiva na análise do domínio e das funcionalidades, com evidências quantitativas obtidas na execução dos testes.")
add_heading("3.2 Levantamento e revisão", 2)
add_plain("A primeira etapa reuniu a monografia anterior, o manual institucional, a documentação do repositório e o código. Cada funcionalidade foi classificada como implementada, condicional, parcial ou futura. Essa classificação foi baseada em evidências de interface, API, regra de negócio, persistência e teste, evitando considerar documentos de intenção como prova de entrega.")
add_heading("3.3 Desenvolvimento incremental", 2)
add_plain("O histórico do repositório indica evolução incremental da aplicação, com mudanças de interface, banco, testes e documentação. A descrição metodológica limita-se ao processo observável. Não se atribui formalmente Scrum ou Kanban ao projeto porque não foram identificadas evidências suficientes de execução completa desses métodos.")
add_heading("3.4 Modelagem", 2)
add_plain("Foram produzidos casos de uso, atividades, sequências, estados, domínio, componentes, implantação e DER. Os modelos foram confrontados com rotas, models, migrations e políticas. Elementos antigos, como pagamento e entregador, foram removidos dos diagramas da versão implementada.")
add_heading("3.5 Implementação", 2)
add_plain("A implementação utiliza Next.js, React e TypeScript na aplicação, PostgreSQL na persistência, Sequelize como ORM e Zod na validação. CSS Modules organizam os estilos. Leaflet, OpenStreetMap e Photon apoiam a busca geográfica. bcrypt protege senhas e sessões assinadas identificam os usuários.")
add_heading("3.6 Validação", 2)
add_plain("A validação combinou testes automatizados, ESLint, build de produção, verificação do schema, conferência de RLS, smoke test transacional e validação do conjunto demonstrativo. Os comandos foram executados em 3 de outubro de 2026 sobre o commit ea44e5e. Propriedades não medidas, como disponibilidade contínua e desempenho sob carga, foram registradas como limitações.")
add_heading("3.7 Rastreabilidade", 2)
add_plain("Foi criada uma matriz relacionando requisitos, interfaces, endpoints, tabelas e testes. Essa matriz sustenta a análise dos resultados e permite identificar áreas que ainda necessitam de evidência dedicada.")


def parse_markdown(path, start_prefix):
    lines = path.read_text(encoding="utf-8").splitlines()
    started = False
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not started:
            if line.startswith(start_prefix):
                started = True
            else:
                i += 1
                continue
        if line.startswith("> **Nota editorial") or line.startswith("> Inserir"):
            i += 1
            continue
        if line.startswith("**Figura X"):
            i += 1
            continue
        if line.startswith("Fonte: Elaborado"):
            i += 1
            continue
        if line.startswith("## "):
            add_heading(line[3:].strip(), 1)
        elif line.startswith("### "):
            add_heading(line[4:].strip(), 2)
        elif line.startswith("#### "):
            add_heading(line[5:].strip(), 3)
        elif line.startswith("**Quadro"):
            add_plain(line.replace("**", ""), WD_ALIGN_PARAGRAPH.CENTER, False, bold=True, size=10, spacing=1)
        elif line.startswith("|") and i + 1 < len(lines) and lines[i+1].startswith("| ---"):
            headers = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            rows=[]
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip().replace("`", "") for c in lines[i].strip("|").split("|")])
                i += 1
            add_table(headers, rows, font_size=8)
            continue
        elif line.startswith("- "):
            add_bullets([line[2:].strip()])
        elif line and not line.startswith(">"):
            add_plain(line.replace("**", "").replace("`", ""))
        i += 1


# Capítulos 4 e 5 já redigidos
parse_markdown(ROOT / "docs" / "Monografia_Parte_7_Capitulo_4_Especificacao.md", "## 4 ")
doc.add_picture(str(use_img), width=Cm(15.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption("Figura 1 - Casos de uso por perfil")
doc.add_picture(str(state_img), width=Cm(15.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption("Figura 2 - Estados da ordem de serviço")

parse_markdown(ROOT / "docs" / "Monografia_Parte_8_Capitulo_5_Modelagem_Arquitetura.md", "## 5 ")
doc.add_picture(str(arch_img), width=Cm(15.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption("Figura 3 - Arquitetura lógica do SmartFix")
doc.add_picture(str(der_img), width=Cm(15.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption("Figura 4 - Grupos de entidades da persistência")

# Capítulo 6
add_heading("6 DESENVOLVIMENTO E IMPLEMENTAÇÃO", 1)
add_heading("6.1 Organização do repositório", 2)
add_plain("O repositório está dividido em documentação, banco e aplicação. A pasta Database reúne tabelas-base, migrations e DER. A pasta docs mantém escopo, requisitos e diagramas. A aplicação Next.js está em smartfix-app, com páginas em app, regras e modelos em src, scripts operacionais e testes automatizados.")
add_table(["Diretório", "Finalidade"], [["Database", "Schema, migrations, DER e documentação do banco."], ["docs", "Escopo, requisitos, diagramas e documentação acadêmica."], ["smartfix-app/app", "Páginas, layouts, componentes e endpoints."], ["smartfix-app/src", "Controllers, serviços, validações, models e tipos."], ["smartfix-app/tests", "Testes automatizados."], ["smartfix-app/scripts", "Migrations, verificações, smoke test e dados demo."]], widths=[4,11])
add_heading("6.2 Interface web", 2)
add_plain("As áreas do cliente e do parceiro possuem layouts próprios e navegação coerente com o papel. Componentes interativos consultam endpoints internos e apresentam estados de carregamento, erro e ausência de dados. CSS Modules isolam estilos das páginas, enquanto globals.css mantém definições globais.")
add_heading("6.3 APIs e servidor", 2)
add_plain("Foram identificados 34 Route Handlers. Eles cobrem autenticação, perfil, endereços, dispositivos, parceiros, serviços, ordens, notificações, localização, avaliações e suporte. Controllers e serviços centralizam a lógica para evitar que a autorização dependa de componentes do navegador.")
add_heading("6.4 Validação e erros", 2)
add_plain("Schemas Zod validam documentos, contatos, endereços, dispositivos, perfil e ações sobre ordens. A resposta da API utiliza formato controlado e evita expor erros de parser ou detalhes internos. Limites de tamanho também são aplicados a descrições, orçamento, comentários e mensagens.")
add_heading("6.5 Persistência", 2)
add_plain("O Sequelize mapeia as tabelas relacionais e participa de transações. As migrations registram evolução do schema, normalização de endereços, criação das ordens, avaliações, serviços e mensagens. O repository de ordens traduz dados do domínio e mantém orçamento, diagnóstico, histórico e avaliação.")
add_heading("6.6 Autenticação e autorização", 2)
add_plain("O login convencional verifica a senha com bcrypt e cria sessão assinada. Google OAuth é opcional. Cada operação protegida recupera o ator da sessão, confere o papel e aplica o identificador como filtro. A administração exige inclusão explícita em ADMIN_USER_IDS.")
add_heading("6.7 Descoberta geográfica", 2)
add_plain("A busca aceita endereço cadastrado ou origem informada, limita o raio aos valores previstos e consulta somente parceiros aprovados. O serviço de geocodificação não inventa coordenadas ausentes. A distância é calculada em linha geográfica e os resultados são ordenados antes da apresentação no mapa Leaflet.")
add_heading("6.8 Fluxo da ordem", 2)
add_plain("A política de ordens concentra ações de orçamento, aprovação, cancelamento, mudança de estado e avaliação. Essa centralização faz com que as mesmas restrições sejam aplicadas independentemente da tela. A persistência registra cada mudança no histórico.")
add_heading("6.9 Interfaces do cliente", 2)
add_plain("O dashboard destaca orçamentos aguardando decisão, reparos prontos e avaliações disponíveis. Endereços e dispositivos possuem operações de manutenção. A busca de assistências combina lista e mapa. A solicitação reúne problema, sintomas e checklist. Ordens, notificações, avaliações, perfil e central de ajuda completam a jornada.")
add_heading("6.10 Interfaces do parceiro e administração", 2)
add_plain("O parceiro consulta um resumo operacional, mantém o catálogo e gerencia ordens. O orçamento é composto por diagnóstico e itens com quantidade e preço. A interface apresenta somente as próximas ações compatíveis com o estado. A administração lista parceiros e controla a aprovação necessária à descoberta.")
add_heading("6.11 Execução", 2)
add_plain("A execução exige Node.js, dependências instaladas, PostgreSQL e variáveis de ambiente. Depois das tabelas-base e migrations, a aplicação pode ser iniciada em desenvolvimento. Antes de publicar, o projeto recomenda db:check, testes, lint e build.")

# Capítulo 7
add_heading("7 TESTES RESULTADOS E DISCUSSÃO", 1)
add_heading("7.1 Ambiente", 2)
add_plain("A validação foi realizada em Windows, com Node.js 24.19.0, npm 11.17.0, Next.js 16.3.4, React 19.2.8, TypeScript 5.9.3 e ESLint 9.39.5. O projeto possuía 16 arquivos de teste, 34 Route Handlers, 23 páginas e oito migrations SQL.")
add_heading("7.2 Testes automatizados", 2)
add_plain("O comando npm test executou 42 casos. Todos foram aprovados, sem falhas, cancelamentos ou itens ignorados. A duração foi de aproximadamente 14 segundos no ambiente utilizado, valor apresentado apenas como registro da execução e não como benchmark.")
add_table(["Indicador", "Resultado"], [["Testes", "42"], ["Aprovados", "42"], ["Falhas", "0"], ["Ignorados", "0"], ["Cancelados", "0"]], widths=[8,7])
add_heading("7.3 Lint e build", 2)
add_plain("O ESLint terminou com código zero e sem diagnósticos. O build de produção compilou a aplicação, executou a verificação TypeScript, coletou dados das páginas, processou a geração estática e finalizou a otimização. O resultado demonstra que a versão analisada pode ser compilada para produção, sem afirmar que serviços opcionais estejam configurados.")
add_heading("7.4 Banco de dados", 2)
add_plain("O db:check confirmou compatibilidade dos modelos principais e RLS ativo nas nove tabelas avaliadas. O smoke test verificou CRUD, propriedade, chaves estrangeiras, cascatas, orçamento, avaliação única e bloqueio de acesso direto pelos papéis anon e authenticated. As fixtures foram revertidas ao final.")
add_heading("7.5 Atendimento dos objetivos", 2)
add_plain("Os resultados confirmam a implementação do fluxo central: cadastro, dispositivos, descoberta, solicitação, diagnóstico, orçamento, aprovação, reparo, conclusão e avaliação. Também foram verificados mecanismos de isolamento, validação e persistência. A rastreabilidade relaciona esses resultados aos requisitos.")
add_heading("7.6 Limitações", 2)
add_plain("A validação não comprova disponibilidade 24 horas, tempo de resposta máximo, capacidade sob carga, acessibilidade integral, segurança absoluta ou compatibilidade com todos os navegadores. Google OAuth e e-mail dependem de configuração. As interfaces completas de atendimento administrativo e do parceiro permanecem parciais.")
add_heading("7.7 Discussão", 2)
add_plain("A principal contribuição técnica está na coerência entre regras, interface e banco. A máquina de estados evita progressões indevidas; as chaves relacionais reforçam propriedade; e os testes cobrem cenários de autorização. Ao mesmo tempo, a separação explícita entre entregue e planejado impede que pagamento e logística sejam avaliados como resultados da versão atual.")

# Capítulo 8
add_heading("8 CONSIDERAÇÕES FINAIS", 1)
add_plain("Este trabalho apresentou a atualização e a validação do SmartFix como plataforma web para organizar o relacionamento entre clientes e assistências técnicas. A análise demonstrou que o sistema atual difere substancialmente do protótipo móvel inicialmente documentado e exigiu revisão de requisitos, diagramas, arquitetura e descrição das interfaces.")
add_plain("O objetivo geral foi atendido dentro do escopo delimitado. A plataforma permite cadastrar contas, endereços e dispositivos, descobrir assistências aprovadas, abrir solicitações, registrar diagnóstico e orçamento, controlar a decisão do cliente e acompanhar as etapas do reparo. Notificações, mensagens e avaliações ampliam a rastreabilidade do relacionamento.")
add_plain("A validação forneceu evidências concretas: 42 testes aprovados, lint e build concluídos, schema compatível, RLS ativo e smoke test transacional bem-sucedido. Essas evidências sustentam os resultados, mas não eliminam as limitações registradas.")
add_heading("8.1 Contribuições", 2)
add_plain("As contribuições incluem a formalização de 47 requisitos funcionais, 41 regras de negócio, 14 requisitos não funcionais sustentados, modelos atualizados, matriz de rastreabilidade e uma arquitetura coerente com o código. O projeto também oferece base executável para evolução acadêmica e técnica.")
add_heading("8.2 Trabalhos futuros", 2)
add_bullets(["integrar pagamento por PIX e cartão com conciliação e repasse;", "implementar coleta, entrega e rastreamento;", "concluir interfaces de atendimento do parceiro e administrador;", "adicionar relatórios e moderação;", "implementar upload seguro de mídias;", "avaliar instalação como PWA;", "realizar testes de carga, acessibilidade e usabilidade com participantes;", "estabelecer monitoramento, backup e evidências de disponibilidade em produção."])
add_heading("8.3 Encerramento", 2)
add_plain("O SmartFix demonstra como princípios de engenharia de software podem transformar um processo informal em um fluxo estruturado. A continuidade do projeto deverá preservar a correspondência entre documentação, código e evidência de validação, especialmente quando novos módulos ampliarem o tratamento de dados e as responsabilidades operacionais.")

# Referências
add_heading("REFERÊNCIAS", 1)
refs = [
    "BRASIL. Lei nº 13.709, de 14 de agosto de 2018. Lei Geral de Proteção de Dados Pessoais. Brasília, DF: Presidência da República, 2018. Disponível em: https://www.planalto.gov.br/ccivil_03/_ato2015-2018/2018/lei/l13709.htm. Acesso em: 3 out. 2026.",
    "NEXT.JS. App Router. 2026a. Disponível em: https://nextjs.org/docs/app. Acesso em: 3 out. 2026.",
    "NEXT.JS. Route Handlers. 2026b. Disponível em: https://nextjs.org/docs/app/getting-started/route-handlers. Acesso em: 3 out. 2026.",
    "POSTGRESQL GLOBAL DEVELOPMENT GROUP. Row Security Policies. PostgreSQL 18 Documentation. 2026. Disponível em: https://www.postgresql.org/docs/18/ddl-rowsecurity.html. Acesso em: 3 out. 2026.",
    "PRESSMAN, Roger S.; MAXIM, Bruce R. Engenharia de software: uma abordagem profissional. 9. ed. Porto Alegre: AMGH, 2021.",
    "REACT TEAM. React v19. 2024. Disponível em: https://react.dev/blog/2024/12/05/react-19. Acesso em: 3 out. 2026.",
    "SOMMERVILLE, Ian. Engenharia de software. 10. ed. São Paulo: Pearson, 2019.",
    "TYPESCRIPT. TypeScript Documentation. 2026. Disponível em: https://www.typescriptlang.org/docs/. Acesso em: 3 out. 2026.",
]
for ref in refs:
    p = add_plain(ref, WD_ALIGN_PARAGRAPH.LEFT, False, size=12, spacing=1)
    p.paragraph_format.space_after = Pt(6)

# Apêndices resumidos
add_heading("APÊNDICE A - REQUISITOS FUNCIONAIS", 1)
rf_rows = []
groups = [
    ("RF-01 a RF-10", "Cadastro, autenticação, sessão, recuperação, OAuth e perfil."),
    ("RF-11 a RF-18", "Endereços e dispositivos."),
    ("RF-19 a RF-23", "Descoberta geográfica, mapa, serviços e avaliações."),
    ("RF-24 a RF-33", "Solicitação, triagem, orçamento, decisão e estados."),
    ("RF-34 a RF-38", "Avaliações, notificações e suporte."),
    ("RF-39 a RF-43", "Dashboard, catálogo, ordens e suporte do parceiro."),
    ("RF-44 a RF-47", "Autorização e operações administrativas."),
]
add_table(["Identificação", "Grupo funcional"], groups, widths=[4,11])
add_plain("O catálogo detalhado, com critérios de aceite e evidências, está preservado na documentação técnica do repositório.")

add_heading("APÊNDICE B - REQUISITOS NÃO FUNCIONAIS", 1)
add_table(["ID", "Propriedade"], [[f"RNF-{i:02d}", desc] for i, desc in enumerate([
    "Hash bcrypt para senhas", "Sessões assinadas", "Autorização por papel e propriedade", "Validação de entrada", "Banco acessado pelo servidor", "RLS e bloqueio público", "Integridade relacional", "Isolamento", "Erros controlados", "Responsividade", "Separação em camadas", "Testes lint e build", "Histórico da ordem", "Falha controlada de integrações"
],1)], widths=[3,12])

add_heading("APÊNDICE C - REGRAS DO CICLO DA ORDEM", 1)
add_table(["Origem", "Ação", "Destino", "Responsável"], [
    ["Novo", "Criar solicitação", "pending", "Cliente"],
    ["pending", "Emitir orçamento", "quoted", "Parceiro"],
    ["pending ou quoted", "Cancelar", "cancelled", "Cliente"],
    ["quoted", "Aprovar", "approved", "Cliente"],
    ["approved", "Iniciar reparo", "in_progress", "Parceiro"],
    ["in_progress", "Aguardar peças", "waiting_parts", "Parceiro"],
    ["waiting_parts", "Retomar reparo", "in_progress", "Parceiro"],
    ["in_progress", "Finalizar reparo", "ready", "Parceiro"],
    ["ready", "Concluir", "completed", "Parceiro"],
], widths=[3.2,5,3.2,3.5])

add_heading("APÊNDICE D - ROTAS PRINCIPAIS DA API", 1)
api_rows = [
    ["/api/auth/*", "Cadastro, login, sessão, logout, recuperação e Google."],
    ["/api/clients/me", "Consulta e atualização do perfil."],
    ["/api/clients/addresses/*", "Endereços e principal."],
    ["/api/clients/devices/*", "Dispositivos."],
    ["/api/partners/nearby", "Busca geográfica."],
    ["/api/partners/[id]/services", "Serviços públicos do parceiro."],
    ["/api/orders/*", "Consulta, criação e atualização das ordens."],
    ["/api/notifications/*", "Consulta e leitura."],
    ["/api/support/*", "Destinatários e mensagens."],
    ["/api/services/*", "Catálogo do parceiro."],
    ["/api/admin/*", "Aprovação e conversas administrativas."],
]
add_table(["Rota", "Finalidade"], api_rows, widths=[6,9])

add_heading("APÊNDICE E - COMANDOS DE EXECUÇÃO E VALIDAÇÃO", 1)
for cmd, desc in [
    ("npm ci", "Instala dependências reproduzíveis."),
    ("npm run db:migrate", "Aplica migrations pendentes."),
    ("npm run db:check", "Confere schema, RLS e acesso do servidor."),
    ("npm run db:smoke", "Executa teste transacional e rollback."),
    ("npm test", "Executa a suíte automatizada."),
    ("npm run lint", "Executa a análise estática."),
    ("npm run build", "Gera e valida o build de produção."),
    ("npm run dev", "Inicia o ambiente de desenvolvimento."),
]:
    add_plain(f"{cmd} - {desc}", WD_ALIGN_PARAGRAPH.LEFT, False, size=12, spacing=1.5)

# Rodapé e numeração
for sec in doc.sections:
    sec.top_margin = Cm(3)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    add_page_number(sec.footer.paragraphs[0])

# Metadados
doc.core_properties.title = f"{TITLE} - Monografia atualizada"
doc.core_properties.subject = "Plataforma web para acompanhamento de reparos de dispositivos eletrônicos"
doc.core_properties.author = "; ".join(AUTHORS)
doc.core_properties.keywords = "SmartFix, engenharia de software, plataforma web, assistência técnica"

doc.save(OUT)
print(OUT)
